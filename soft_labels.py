"""Strict ratings joins onto the existing, immutable image population and split."""
import csv
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd


TARGET_COLUMNS = [f'target_grade_{grade}' for grade in range(1, 6)]


def read_table(path, sheet='Grades'):
    if Path(path).suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        with Path(path).open('rb') as handle:
            workbook = load_workbook(handle, read_only=True, data_only=False)
            try:
                if sheet not in workbook.sheetnames:
                    raise ValueError(f'{path}: missing worksheet {sheet!r}')
                rows = list(workbook[sheet].values)
            finally:
                workbook.close()
        columns = list(rows[0])
        if any(not isinstance(c, str) or not c.strip() for c in columns) or len(columns) != len(set(columns)):
            raise ValueError(f'{path}: missing or duplicate column names')
        return pd.DataFrame(rows[1:], columns=columns).fillna('')
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        columns = next(csv.reader(handle))
    if len(columns) != len(set(columns)):
        raise ValueError(f'{path}: duplicate column names')
    return pd.read_csv(path, keep_default_na=False)


def image_key(value, filename=False):
    value = str(value).strip().replace('\\', '/')
    if filename:
        path = Path(value)
        return path.stem if path.suffix.lower() in {'.png', '.tif', '.tiff', '.jpg', '.jpeg', '.bmp', '.gif', '.webp'} else path.name
    return value


def validate_manifest(df, check_files=True):
    required = ['image_id', 'image_path', 'patient_id', 'project', 'project_site', 'split',
                'previousStudent_grade', 'shane_grade']
    if df.empty or not set(required).issubset(df.columns):
        raise ValueError('Use the existing shared manifest from prepare_inputs.py')
    for column in required:
        if df[column].isna().any() or df[column].astype(str).str.strip().eq('').any():
            raise ValueError(f'Manifest has missing {column}')
    if df.image_id.duplicated().any() or df.image_path.map(lambda p: str(Path(p).resolve()).casefold()).duplicated().any():
        raise ValueError('Manifest contains duplicate image IDs or ambiguous image paths')
    if df.image_id.map(image_key).duplicated().any():
        raise ValueError('Manifest contains duplicate normalized image IDs')
    if not df.split.isin(['train', 'validation', 'test']).all():
        raise ValueError('Unknown manifest split')
    if df.groupby('patient_id').split.nunique().max() != 1:
        raise ValueError('Patient leakage in manifest')
    held_out = df.project.str.casefold().eq('visus-nhp')
    if not held_out.any() or df.loc[held_out, 'split'].ne('test').any():
        raise ValueError('Entire VISUS-NHP trial must remain held out')
    if not all(df.split.eq(mode).any() for mode in ('train', 'validation')):
        raise ValueError('Empty train or validation split')
    for column in ('previousStudent_grade', 'shane_grade'):
        if not pd.to_numeric(df[column], errors='coerce').isin(range(1, 6)).all():
            raise ValueError(f'Invalid manifest {column}')
    if not all(image_key(path, filename=True) == image_id for path, image_id in zip(df.image_path, df.image_id)):
        raise ValueError('Manifest image ID/path mismatch (retain full variant identifiers)')
    if check_files and not df.image_path.map(lambda p: Path(p).is_file()).all():
        raise ValueError('Manifest image files are missing')


def construct_targets(df, columns):
    if not columns or len(columns) != len(set(columns)) or not set(columns).issubset(df.columns):
        raise ValueError('Select unique, existing grader columns')
    if any(not column.endswith('_grade') or column in TARGET_COLUMNS for column in columns):
        raise ValueError('Grader columns must end in _grade; target columns cannot be graders')
    result = df.copy()
    numeric = []
    for column in columns:
        raw = result[column]
        missing = raw.isna() | raw.astype(str).str.strip().eq('')
        values = pd.to_numeric(raw.where(~missing), errors='coerce')
        invalid = ~missing & (~values.isin(range(1, 6)) | raw.map(lambda value: isinstance(value, (bool, np.bool_))))
        if invalid.any():
            index = invalid[invalid].index[0]
            raise ValueError(f'Invalid rating: image {result.loc[index, "image_id"]!r}, column {column!r}, value {raw.loc[index]!r}; expected integer 1-5 or blank')
        numeric.append(values.to_numpy(dtype=float))
    ratings = np.stack(numeric, axis=1)
    counts = np.isfinite(ratings).sum(axis=1)
    targets = np.stack([(ratings == grade).sum(axis=1) for grade in range(1, 6)], axis=1)
    targets = np.divide(targets, counts[:, None], out=np.zeros_like(targets, dtype=float), where=counts[:, None] != 0)
    result[TARGET_COLUMNS] = targets
    result['rating_count'] = counts
    result['has_ratings'] = counts > 0
    return result


def load_soft_manifest(manifest, ratings_path, columns=None, check_files=True, sheet='Grades', exclude_known_unmatched=False):
    base = read_table(manifest)
    validate_manifest(base, check_files=check_files)
    ratings = read_table(ratings_path, sheet=sheet)
    if 'image_id' in ratings and 'image_filename' in ratings:
        if not (ratings.image_id.map(image_key) == ratings.image_filename.map(lambda x: image_key(x, True))).all():
            raise ValueError('Ambiguous image_id/image_filename join')
    key = 'image_id' if 'image_id' in ratings else 'image_filename'
    if key not in ratings:
        raise ValueError('Ratings require image_id or image_filename')
    ids = ratings[key].map(lambda x: image_key(x, key == 'image_filename'))
    if ids.eq('').any() or ids.duplicated().any():
        raise ValueError('Ratings contain empty or duplicate image IDs')
    columns = list(columns) if columns else [c for c in ratings if c.endswith('_grade') and c not in TARGET_COLUMNS]
    ratings = ratings.copy()
    ratings['image_id'] = ids
    construct_targets(ratings, columns)
    unknown = sorted(set(ids) - set(base.image_id))
    if unknown and exclude_known_unmatched:
        report_path = Path(manifest).parent / 'split_report.json'
        report = json.loads(report_path.read_text(encoding='utf-8'))
        known = set(report.get('shane_only_ids', [])) | set(report.get('previousStudent_only_ids', []))
        if not set(unknown).issubset(known):
            raise ValueError(f'Unmatched ratings not in saved split exclusions: {sorted(set(unknown) - known)[:5]}')
        print(f'Explicitly excluding {len(unknown)} previously unmatched ratings IDs from {report_path}: {unknown}')
        ratings = ratings.loc[~ids.isin(unknown)].copy()
        ids = ids.loc[ratings.index]
        unknown = []
    if unknown:
        raise ValueError(f'{len(unknown)} ratings IDs outside existing manifest; no images may be added silently: {unknown[:5]}')
    aligned = ratings.set_index('image_id').reindex(base.image_id)
    for column in set(base.columns) & set(ratings.columns) - {'image_id', 'image_filename'}:
        supplied = aligned[column].notna() & aligned[column].astype(str).str.strip().ne('')
        expected = base.set_index('image_id')[column]
        if column.endswith('_grade'):
            differs = pd.to_numeric(aligned[column], errors='coerce') != pd.to_numeric(expected, errors='coerce')
        else:
            differs = aligned[column].astype(str) != expected.astype(str)
        if (supplied & differs).any():
            raise ValueError(f'Ratings conflict with existing manifest {column}; original grades/metadata/splits must be preserved')
    for column in columns:
        base[column] = aligned[column].to_numpy()
    result = construct_targets(base, columns)
    for column in columns:
        available = result[column].notna() & result[column].astype(str).str.strip().ne('')
        print(f'{column}: available ratings by split {result.loc[available].groupby("split").size().to_dict()}')
    report = {mode: int((result.split.eq(mode) & ~result.has_ratings).sum()) for mode in ('train', 'validation', 'test')}
    print(f'Images with no selected ratings (excluded from loss/evaluation, retained in snapshot): {report}')
    for mode in ('train', 'validation'):
        if not (result.split.eq(mode) & result.has_ratings).any():
            raise ValueError(f'No rated images in {mode}')
    return result, columns, report


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()
